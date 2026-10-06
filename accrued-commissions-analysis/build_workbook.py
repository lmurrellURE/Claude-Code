"""Build the Accrued Commissions / Unclaimed Property analysis workbook.

Inputs : accrued_rows.csv, current_month_unbooked.csv (extracted from United_Cash_Requirements.xlsx)
Output : Accrued_Commissions_Unclaimed_Property_Analysis.xlsx
"""
import csv, datetime, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1] if len(sys.argv) > 1 else "Accrued_Commissions_Unclaimed_Property_Analysis.xlsx"
AS_OF = datetime.date(2026, 10, 6)

F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True)
f_input = Font(name=F, size=10, color="0000FF")          # blue = hardcoded input
f_link = Font(name=F, size=10, color="008000")           # green = cross-sheet link
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
fill_hdr = PatternFill("solid", fgColor="1F3864")
fill_in = PatternFill("solid", fgColor="FFFF00")          # yellow = user fills in
fill_sub = PatternFill("solid", fgColor="D9E1F2")
fill_warn = PatternFill("solid", fgColor="FCE4D6")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
CUR = '$#,##0.00;($#,##0.00);-'
DATE = 'mm/dd/yyyy'

def hdr(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row, i, lab); c.font = f_hdr; c.fill = fill_hdr
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = box
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[row].height = 42

def put(ws, ref, val, font=f_norm, fmt=None, fill=None, wrap=False):
    c = ws[ref]; c.value = val; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if wrap: c.alignment = Alignment(wrap_text=True, vertical="top")
    return c

# ---------------------------------------------------------------- data
rows = [r for r in csv.DictReader(open("accrued_rows.csv")) if r["booked"] == "False" and r["amount"]]
for r in rows:
    r["amt"] = float(r["amount"])
    r["d"] = datetime.date.fromisoformat(r["date"]) if r["date"] else None
cur = list(csv.DictReader(open("current_month_unbooked.csv")))

ENTITIES = [  # entity, state (assumption flag), note
    ("Dallas", "TX", "Per user: Texas office."),
    ("Houston", "TX", "Per user: Texas office."),
    ("TUR", "TX", "Texas United Realty - per user: Texas office."),
    ("Chicago", "IL", "Per user: Illinois office."),
    ("Philly", "PA", "Per user: Philadelphia office."),
    ("Gallery", "FL", "ASSUMPTION - user listed Florida and Alabama offices without naming tabs; Gallery banks at SouthState (a Florida-based bank). CONFIRM."),
    ("Leading Edge", "AL", "ASSUMPTION - user listed Florida and Alabama offices without naming tabs; Leading Edge banks at Regions/Cadence (Alabama-based banks). CONFIRM."),
    ("DC", "DC", "Tab exists in workbook with $0 accrued balance; user did not list DC. No DC rules researched."),
]
STATE_ROWS = [
    # state, name, short_period_years, short_basis, long_period_years, long_basis, report_due, cutoff, next_cutoff, due_diligence, vda, verification, sources
    ("TX", "Texas", 1, "Unclaimed 'wages' - Tex. Prop. Code 72.1015 (1 year). 'Wages' takes its meaning from Labor Code 61.001, which covers compensation owed by an EMPLOYER to an EMPLOYEE and excludes independent contractors. Independent-contractor agent commissions therefore likely fall under the general rule, not the 1-year rule - confirm with counsel / Comptroller.",
     3, "General personal property - Tex. Prop. Code 72.101(a): presumed abandoned after more than 3 years with owner location unknown and no claim/act of ownership.",
     "July 1 (Tex. Prop. Code 74.101)", "March 1 (property held on March 1 is reported the following July 1)", datetime.date(2027, 3, 1),
     "Notice to owner by mail/e-mail at least 60 days before delivery for property over $250 (Tex. Prop. Code 74.1011).",
     "Comptroller offers a Voluntary Disclosure Agreement with waiver of penalty and interest; not available once an investigation has started or (per secondary sources) to holders who have filed Texas reports before.",
     "Dormancy periods, report date, and due-diligence rule confirmed from search summaries of the statute text; statute pages themselves were blocked by this session's network policy. No business-to-business exemption was found for Texas.",
     "https://statutes.capitol.texas.gov/docs/PR/htm/PR.72.htm ; https://statutes.capitol.texas.gov/Docs/PR/htm/PR.74.htm ; https://texas.public.law/statutes/tex._prop._code_section_72.1015 ; https://texas.public.law/statutes/tex._labor_code_section_61.001 ; https://www.comptroller.texas.gov/taxes/publications/96-576.php"),
    ("IL", "Illinois", 1, "765 ILCS 1026/15-201: 'wages, commissions, bonuses, or reimbursements to which an employee is entitled, or other compensation for personal services' - 1 year after payable. The 'other compensation for personal services' language is broader than Texas and plausibly reaches independent-contractor commissions.",
     3, "765 ILCS 1026/15-201 catch-all: all other property 3 years after the obligation to pay arises (per secondary sources).",
     "November 1 for financial, insurance and government holders (confirmed). Deadline for ordinary business holders NOT confirmed from the statute in this session - the prior Illinois act used May 1 with a December 31 cutoff. CONFIRM.", "June 30 for the November 1 filers; business-holder cutoff NOT confirmed.", datetime.date(2027, 6, 30),
     "Due-diligence letter by first-class mail for property over $50, sent 60 to 365 days before the report (per Illinois Treasurer guidance summary).",
     "Not researched in this session.",
     "Illinois eliminated its business-to-business exemption when it adopted RUUPA (effective 1/1/2018). Chicago's prior-month accrued balance is $0, so Illinois exposure today is only the current-month unbooked deposits.",
     "https://www.ilga.gov/legislation/ilcs/documents/076510260K15-201.htm ; https://icash.illinoistreasurer.gov/app/faq ; https://www.duanemorris.com/articles/illinois_slides_backward_unclaimed_property_law_0917.html"),
    ("PA", "Pennsylvania", 2, "Wages / commissions - 2 years per the PA Treasury Dormancy Matrix as summarized in search results (cited there as 72 P.S. 1301.5). NOT verified from the statute text in this session.",
     3, "General dormancy period 3 years for most property types (72 P.S. 1301.1 et seq.) per secondary sources.",
     "April 15", "December 31 (per secondary sources; NOT verified from statute in this session)", datetime.date(2026, 12, 31),
     "Written notice by first-class mail 60 to 120 days before the April 15 deadline for property of $50 or more with a usable address.",
     "PA's formal amnesty program ended 10/31/2010. Interest of 12% per year from the date the property was due is cited by secondary sources. Ask PA Treasury about current voluntary compliance terms.",
     "All PA facts come from secondary summaries (PA Treasury Dormancy Matrix as quoted, Sovos, UPPO). Statute and Treasury pages were blocked in this session.",
     "https://www.patreasury.gov/pdf/unclaimed-property/Dormancy-Matrix.pdf ; https://www.patreasury.gov/unclaimed-property/holders/ ; https://sovos.com/tax-reporting/unclaimed-property-laws-by-state/pennsylvania/"),
    ("FL", "Florida", 1, "Fla. Stat. 717.115: unpaid WAGES owing in the ordinary course of business unclaimed for more than 1 year after becoming payable. Whether independent-contractor commissions are 'wages' here is not settled in the sources found - if not, the 5-year general rule applies.",
     5, "Fla. Stat. 717.102(1): general 5-year dormancy period for other intangible property.",
     "Before May 1, covering the preceding calendar year (Fla. Stat. 717.117)", "December 31", datetime.date(2026, 12, 31),
     "Report must list owners with property of $50 or more; due-diligence notice requirement in 717.117 (threshold and timing not verified from statute text in this session).",
     "Florida has a Voluntary Disclosure program: no penalties, 5-year look-back instead of 10, 2-year audit window after acceptance (per secondary summary of Florida DFS program).",
     "Dormancy periods and report date confirmed from search summaries of the statute; statute pages blocked. Gallery = Florida is an ASSUMPTION.",
     "https://www.flsenate.gov/Laws/Statutes/2024/717.115 ; https://www.flsenate.gov/Laws/Statutes/2024/717.102 ; https://flsenate.gov/Laws/Statutes/2024/717.117 ; https://www.fltreasurehunt.gov/Holder.jsp"),
    ("AL", "Alabama", 1, "Ala. Code 35-12-72: 'wages or other compensation for personal services' presumed abandoned 1 year after the compensation becomes payable. 'Other compensation for personal services' plausibly reaches independent-contractor commissions.",
     3, "Ala. Code 35-12-72: all other property 3 years after the owner's right to demand it or the obligation to pay arises, whichever is first.",
     "November 1, covering the 12 months preceding July 1", "June 30", datetime.date(2027, 6, 30),
     "Aggregate/reporting threshold $50 per secondary sources; due-diligence notice rule (35-12-77) not verified in this session.",
     "Not researched in this session.",
     "Dormancy periods confirmed from search summaries of the Code; Code pages blocked. Leading Edge = Alabama is an ASSUMPTION.",
     "https://law.justia.com/codes/alabama/2023/title-35/chapter-12/article-2/article-2a/section-35-12-72 ; https://treasury.alabama.gov/alabama-treasury-law/unclaimed-property-law/ ; https://unclaimed.org/reporting/alabama/"),
    ("MO", "Missouri (HQ / possible state of formation)", 3, "RSMo 447.536 as amended by HB 1075 (2014): payroll/wages dormancy reduced from 5 to 3 years effective 1/1/2015.",
     5, "RSMo 447.536: general 5-year dormancy period for most property types.",
     "November 1", "June 30", datetime.date(2027, 6, 30),
     "Due diligence letter required when the owner is due $50 or more (Missouri Treasurer guidance).",
     "Not researched in this session.",
     "Missouri matters only under the SECOND priority rule (Texas v. New Jersey, 379 U.S. 674 (1965)): property whose owner address is unknown goes to the holder's state of incorporation/formation - which is the entity's formation state, NOT its headquarters city. Missouri also has a business-to-business exclusion (HB 1075, 2014) for checks/credits owed to a business entity in the ordinary course of business.",
     "https://revisor.mo.gov/main/OneSection.aspx?section=447.536 ; https://treasurer.mo.gov/UCP/ReportingUnclaimedProperty.aspx ; https://ryan.com/about-ryan/news-and-insights/2014/missouri-legislation-adds-unclaimed-property-business-to-business-exclusions-and-other-changes/"),
    ("DC", "District of Columbia", None, "Not researched - DC accrued balance is $0.", None, "Not researched.", "", "", None, "", "", "Not researched.", ""),
]

wb = Workbook()

# ================================================================ State Rules
sr = wb.active; sr.title = "State Rules"
put(sr, "A1", "Unclaimed property rules by state - dormancy periods applied in this analysis", f_title)
put(sr, "A2", "Short period = the period for wages / compensation for personal services. Long period = the general catch-all period. Which one applies to independent-contractor agent commissions is a legal question that differs by state (see basis columns). Verification column says what could and could not be confirmed from primary sources in this session.", wrap=True)
sr.merge_cells("A2:N2"); sr.row_dimensions[2].height = 45
labels = ["State", "Name", "Short period (yrs)", "Short period basis", "Long period (yrs)", "Long period basis", "Report due", "Report cutoff ('as of')", "Next cutoff date", "Due diligence to owner", "Voluntary disclosure / penalties", "Verification status", "Sources"]
hdr(sr, 4, labels, [7, 18, 9, 55, 9, 45, 28, 28, 12, 40, 45, 45, 70])
for i, s in enumerate(STATE_ROWS, 5):
    vals = list(s)
    for j, v in enumerate(vals, 1):
        c = sr.cell(i, j, v); c.font = f_input if j in (3, 5, 9) else f_norm
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = box
        if j == 9 and v: c.number_format = DATE
    sr.row_dimensions[i].height = 150
SR_FIRST, SR_LAST = 5, 4 + len(STATE_ROWS)
put(sr, f"A{SR_LAST+2}", "Priority rules (Texas v. New Jersey, 379 U.S. 674 (1965)): (1) the state of the owner's last-known address on the holder's books; (2) if no address, the holder's state of incorporation/formation. Private retention of amounts owed to others ('private escheat') is not permitted by any state; an unknown owner does not make the money the company's.", wrap=True)
sr.merge_cells(f"A{SR_LAST+2}:N{SR_LAST+2}"); sr.row_dimensions[SR_LAST+2].height = 45
put(sr, f"A{SR_LAST+3}", "Blue = hardcoded inputs you can change. All dormancy periods are expressed in years and are applied from the DEPOSIT DATE in the cash-requirements workbook as a proxy for the date the amount became payable - the true trigger is when the agent's share became payable, which may be earlier (closing date) or later.", wrap=True)
sr.merge_cells(f"A{SR_LAST+3}:N{SR_LAST+3}"); sr.row_dimensions[SR_LAST+3].height = 40

# ================================================================ Summary
sm = wb.create_sheet("Summary", 0)
put(sm, "A1", "Accrued Commissions - Unclaimed Property Exposure by Entity", f_title)
put(sm, "A2", "Source: United_Cash_Requirements.xlsx, each entity tab, 'Previous Month's Accrued Commissions' block (deposited checks not yet booked in Sage) plus the current-month 'Deposited checks' block. Extracted 10/6/2026; every entity total ties to the tab's 'outstanding amount' cell.", wrap=True)
sm.merge_cells("A2:P2"); sm.row_dimensions[2].height = 32
put(sm, "A3", "As-of date"); put(sm, "B3", AS_OF, f_input, DATE, fill_in)
sm["B3"].comment = Comment("Input. Change to re-age every item. 10/6/2026 = date the workbook was provided (Summary!A19 of the source file is =TODAY()).", "Analysis")
put(sm, "A4", "Legend: blue = input, black = formula, green = link to another sheet, yellow fill = cells for you to confirm or fill in.", wrap=False)

labels = ["Entity", "State (confirm)", "Prior-month accrued (unbooked deposits)", "Current-month accrued (Sept/Oct 2026 deposits)", "Total accrued per workbook", "Items (count)",
          "0-1 yrs", "1-2 yrs", "2-3 yrs", "3-5 yrs", "5+ yrs",
          "Short period (yrs)", "Long period (yrs)",
          "Dormant today - SHORT period test", "Dormant today - LONG period test",
          "Dormant at next report cutoff - SHORT", "Dormant at next report cutoff - LONG",
          "Oldest item", "Items with memo flag (not ours / unknown / hold)", "$ with memo flag", "Items at recurring flat amounts (possible brokerage fee)", "$ at recurring flat amounts", "Entity-state note"]
hdr(sm, 6, labels, [14, 9, 16, 16, 16, 8, 13, 13, 13, 13, 13, 9, 9, 16, 16, 16, 16, 12, 12, 14, 14, 14, 60])
R0 = 7
for i, (ent, st, note) in enumerate(ENTITIES):
    r = R0 + i
    put(sm, f"A{r}", ent, f_bold)
    put(sm, f"B{r}", st, f_input, fill=fill_in)
    put(sm, f"C{r}", f'=SUMIFS(Detail!$H:$H,Detail!$A:$A,$A{r})', f_link, CUR)
    put(sm, f"D{r}", f"=SUMIFS('Current Month'!$F:$F,'Current Month'!$A:$A,$A{r})", f_link, CUR)
    put(sm, f"E{r}", f"=C{r}+D{r}", fmt=CUR)
    put(sm, f"F{r}", f'=COUNTIFS(Detail!$A:$A,$A{r},Detail!$H:$H,"<>")')
    for j, b in enumerate(["0-1 yrs", "1-2 yrs", "2-3 yrs", "3-5 yrs", "5+ yrs"]):
        put(sm, f"{L(7+j)}{r}", f'=SUMIFS(Detail!$H:$H,Detail!$A:$A,$A{r},Detail!$L:$L,"{b}")', f_link, CUR)
    put(sm, f"L{r}", f"=IFERROR(INDEX('State Rules'!$C${SR_FIRST}:$C${SR_LAST},MATCH($B{r},'State Rules'!$A${SR_FIRST}:$A${SR_LAST},0)),\"\")", f_link)
    put(sm, f"M{r}", f"=IFERROR(INDEX('State Rules'!$E${SR_FIRST}:$E${SR_LAST},MATCH($B{r},'State Rules'!$A${SR_FIRST}:$A${SR_LAST},0)),\"\")", f_link)
    put(sm, f"N{r}", f'=SUMIFS(Detail!$H:$H,Detail!$A:$A,$A{r},Detail!$O:$O,"<="&$B$3)', f_link, CUR)
    put(sm, f"O{r}", f'=SUMIFS(Detail!$H:$H,Detail!$A:$A,$A{r},Detail!$P:$P,"<="&$B$3)', f_link, CUR)
    put(sm, f"P{r}", f'=SUMIFS(Detail!$H:$H,Detail!$A:$A,$A{r},Detail!$Q:$Q,"Yes")', f_link, CUR)
    put(sm, f"Q{r}", f'=SUMIFS(Detail!$H:$H,Detail!$A:$A,$A{r},Detail!$R:$R,"Yes")', f_link, CUR)
    put(sm, f"R{r}", f'=IF(F{r}=0,"",_xlfn.MINIFS(Detail!$E$7:$E$400,Detail!$A$7:$A$400,$A{r}))', f_link, DATE)
    put(sm, f"S{r}", f'=COUNTIFS(Detail!$A:$A,$A{r},Detail!$M:$M,"Memo flag*")')
    put(sm, f"T{r}", f'=SUMIFS(Detail!$H:$H,Detail!$A:$A,$A{r},Detail!$M:$M,"Memo flag*")', f_link, CUR)
    put(sm, f"U{r}", f'=COUNTIFS(Detail!$A:$A,$A{r},Detail!$M:$M,"Recurring flat*")')
    put(sm, f"V{r}", f'=SUMIFS(Detail!$H:$H,Detail!$A:$A,$A{r},Detail!$M:$M,"Recurring flat*")', f_link, CUR)
    put(sm, f"W{r}", note, wrap=True)
    for c in range(1, 24): sm.cell(r, c).border = box
RT = R0 + len(ENTITIES)
put(sm, f"A{RT}", "TOTAL", f_bold)
for col in "CDEFGHIJKNOPQTV":
    put(sm, f"{col}{RT}", f"=SUM({col}{R0}:{col}{RT-1})", f_bold, CUR if col not in "FSU" else None)
for col in "SU":
    put(sm, f"{col}{RT}", f"=SUM({col}{R0}:{col}{RT-1})", f_bold)
for c in range(1, 24): sm.cell(RT, c).border = box; sm.cell(RT, c).fill = fill_sub
sm.freeze_panes = "C7"

n = RT + 2
notes = [
    ("How to read this", f_bold),
    ("'Dormant today - SHORT period test' = unbooked deposits older than the state's wages/compensation period as of the as-of date. This is the UPPER BOUND of what could already be reportable, and only if 100% of each check is owed to someone else (an agent or the payer). It is NOT the amount to remit.", f_norm),
    ("'Dormant today - LONG period test' = older than the state's general catch-all period. For Texas (the three largest balances) this is the more likely test for independent-contractor commissions, because Texas ties 'wages' to the Labor Code definition that excludes independent contractors.", f_norm),
    ("The amount actually reportable = (agent share + any amount belonging to the payer) of each dormant item. The brokerage's own fee / E&O / company-dollar share of a check is the company's revenue and is never unclaimed property. The source workbook has 'E&O', 'FEE' and 'Commission' columns for this split, but they are EMPTY for every item, so the split is unknown today. See the Detail tab's yellow columns.", f_norm),
    ("Current-month accrued (column D) is all September/October 2026 deposits in normal matching workflow. It is included so the total ties to the cash-requirements Summary, but none of it is near any dormancy period.", f_norm),
    ("Nothing here is legal advice. Dormancy periods, report dates and due-diligence rules on the State Rules tab were sourced from statute summaries and state guidance found by web search; the statute pages themselves could not be opened from this session. Have counsel or your unclaimed-property advisor confirm before filing.", f_norm),
]
for i, (t, f) in enumerate(notes):
    put(sm, f"A{n+i}", t, f, wrap=True); sm.merge_cells(f"A{n+i}:P{n+i}"); sm.row_dimensions[n+i].height = 16 if f is f_bold else 44

# ================================================================ Detail
dt = wb.create_sheet("Detail", 1)
put(dt, "A1", "Unbooked deposited checks - item detail (prior-month accrued commissions block of each entity tab)", f_title)
put(dt, "A2", "Columns A-J are copied from the source workbook. K-R are formulas. S-X (yellow) are for you to fill in from Paperless Pipeline / Sage: once the agent share is entered, column Y shows the dollar amount that is both dormant (long test) and owed to someone else.", wrap=True)
dt.merge_cells("A2:Y2"); dt.row_dimensions[2].height = 32
put(dt, "A3", "Classification heuristics (col M) are only hints: 'Recurring flat amount' = $644 / $595 / $495 / $49 (amounts that repeat across unrelated payers and look like a flat transaction or E&O fee the company keeps - CONFIRM). 'Memo flag' = memo says not ours / unknown agent / duplicate / legal hold. Everything else = needs transaction match.", wrap=True)
dt.merge_cells("A3:Y3"); dt.row_dimensions[3].height = 32
put(dt, "A4", "As-of date (linked)"); put(dt, "B4", "=Summary!$B$3", f_link, DATE)
labels = ["Entity", "State", "Bank / block", "Source row", "Deposit date", "Check #", "Payer / check writer", "Amount", "Memo (source)", "Booked in Sage?",
          "Age (yrs)", "Age bucket", "Classification hint",
          "Short period (yrs)", "Dormant date - SHORT", "Dormant date - LONG", "Dormant at next cutoff - SHORT?", "Dormant at next cutoff - LONG?",
          "Agent / owner name (fill in)", "Owner last-known state (fill in)", "Brokerage fee / E&O share $ (fill in)", "Agent share $ (fill in)", "Refund to payer $ (fill in)", "Resolution / status (fill in)",
          "Reportable if dormant (LONG test) = agent + refund share"]
hdr(dt, 6, labels, [13, 7, 13, 8, 11, 14, 32, 12, 55, 8, 8, 9, 30, 8, 12, 12, 11, 11, 22, 10, 13, 13, 13, 28, 16])
FLAT = {644.0, 595.0, 495.0, 49.0}
FLAGS = ["not for this", "not in pp", "not sure", "unsure", "duplicate", "couldn't locate", "cant locate", "can't locate", "fraud", "on hold", "law enforcement", "unidentified", "possibly", "believe this transaction", "shortage", "supposed to be", "review check stub", "no memo", "?"]
def classify(r):
    memo = (r["memo"] or "").lower()
    if any(k in memo for k in FLAGS): return "Memo flag - not ours / unknown / hold: confirm"
    if round(r["amt"], 2) in FLAT: return "Recurring flat amount - possible brokerage fee: confirm"
    return "Needs transaction match"
ent_order = {e[0]: i for i, e in enumerate(ENTITIES)}
rows.sort(key=lambda r: (ent_order.get(r["entity"], 99), r["d"] or datetime.date(1900, 1, 1)))
D0 = 7
for i, r in enumerate(rows):
    rr = D0 + i
    put(dt, f"A{rr}", r["entity"])
    put(dt, f"B{rr}", f"=IFERROR(INDEX(Summary!$B${R0}:$B${RT-1},MATCH($A{rr},Summary!$A${R0}:$A${RT-1},0)),\"\")", f_link)
    put(dt, f"C{rr}", r["bank"]); put(dt, f"D{rr}", int(r["row"]))
    put(dt, f"E{rr}", r["d"], fmt=DATE); put(dt, f"F{rr}", r["check"])
    put(dt, f"G{rr}", r["payer"]); put(dt, f"H{rr}", r["amt"], fmt=CUR)
    put(dt, f"I{rr}", r["memo"], wrap=False); put(dt, f"J{rr}", "No")
    put(dt, f"K{rr}", f"=IF(E{rr}=\"\",\"\",($B$4-E{rr})/365.25)", fmt="0.00")
    put(dt, f"L{rr}", f'=IF(K{rr}="","no date",IF(K{rr}<1,"0-1 yrs",IF(K{rr}<2,"1-2 yrs",IF(K{rr}<3,"2-3 yrs",IF(K{rr}<5,"3-5 yrs","5+ yrs")))))')
    put(dt, f"M{rr}", classify(r))
    put(dt, f"N{rr}", f"=IFERROR(INDEX('State Rules'!$C${SR_FIRST}:$C${SR_LAST},MATCH($B{rr},'State Rules'!$A${SR_FIRST}:$A${SR_LAST},0)),\"\")", f_link)
    put(dt, f"O{rr}", f'=IF(OR(N{rr}="",E{rr}=""),"",EDATE(E{rr},12*N{rr}))', fmt=DATE)
    put(dt, f"P{rr}", f"=IF(E{rr}=\"\",\"\",IFERROR(EDATE(E{rr},12*INDEX('State Rules'!$E${SR_FIRST}:$E${SR_LAST},MATCH($B{rr},'State Rules'!$A${SR_FIRST}:$A${SR_LAST},0))),\"\"))", fmt=DATE)
    nxt = f"INDEX('State Rules'!$I${SR_FIRST}:$I${SR_LAST},MATCH($B{rr},'State Rules'!$A${SR_FIRST}:$A${SR_LAST},0))"
    put(dt, f"Q{rr}", f'=IF(O{rr}="","",IFERROR(IF(O{rr}<={nxt},"Yes","No"),""))')
    put(dt, f"R{rr}", f'=IF(P{rr}="","",IFERROR(IF(P{rr}<={nxt},"Yes","No"),""))')
    for col in "STUVWX": put(dt, f"{col}{rr}", None, f_input, fill=fill_in)
    for col in "UVW": dt[f"{col}{rr}"].number_format = CUR
    put(dt, f"Y{rr}", f'=IF(P{rr}="","",IF(P{rr}<=$B$4,N(V{rr})+N(W{rr}),0))', fmt=CUR)
    for c in range(1, 26): dt.cell(rr, c).border = box
    if "Memo flag" in dt[f"M{rr}"].value: dt[f"M{rr}"].fill = fill_warn
DT_LAST = D0 + len(rows) - 1
put(dt, f"A{DT_LAST+1}", "TOTAL", f_bold)
put(dt, f"H{DT_LAST+1}", f"=SUM(H{D0}:H{DT_LAST})", f_bold, CUR)
for col in "UVWY": put(dt, f"{col}{DT_LAST+1}", f"=SUM({col}{D0}:{col}{DT_LAST})", f_bold, CUR)
dt.freeze_panes = "E7"
dt.auto_filter.ref = f"A6:Y{DT_LAST}"

# ================================================================ Current Month
cm = wb.create_sheet("Current Month", 2)
put(cm, "A1", "Current-month unbooked deposits ('Deposited checks' block, no 'x' in Sage column) - all September/October 2026", f_title)
put(cm, "A2", "Listed only so the Summary ties to the cash-requirements workbook. These are in the normal matching cycle and are not an unclaimed-property concern today.", wrap=True)
cm.merge_cells("A2:G2")
hdr(cm, 4, ["Entity", "Bank / block", "Source row", "Deposit date", "Check #", "Amount"], [14, 14, 9, 12, 16, 14])
for i, r in enumerate(cur):
    rr = 5 + i
    put(cm, f"A{rr}", r["entity"]); put(cm, f"B{rr}", r["bank"]); put(cm, f"C{rr}", int(r["row"]))
    put(cm, f"D{rr}", datetime.date.fromisoformat(r["date"]) if r["date"] else None, fmt=DATE)
    put(cm, f"E{rr}", r["check"]); put(cm, f"F{rr}", float(r["amount"]), fmt=CUR)
    for c in range(1, 7): cm.cell(rr, c).border = box
put(cm, f"A{5+len(cur)}", "TOTAL", f_bold); put(cm, f"F{5+len(cur)}", f"=SUM(F5:F{4+len(cur)})", f_bold, CUR)

# ================================================================ Gaps & Next Steps
gp = wb.create_sheet("Gaps & Next Steps", 3)
put(gp, "A1", "What is missing, and what to do next", f_title)
hdr(gp, 3, ["#", "Item", "Why it matters", "Who / where"], [4, 55, 80, 30])
gaps = [
    ("DATA GAP: the fee vs. agent split of every check is unknown", "The source workbook has E&O / FEE / Commission columns but they are empty for all 181 items. Only the agent's share (and any amount that belongs to the payer) can ever be unclaimed property. The company's own fee is revenue. Until the split is known, every dollar figure on the Summary is an upper bound, not an amount to remit.", "Match each check to its Paperless Pipeline transaction; fill Detail cols U-W."),
    ("DATA GAP: owner identity and last-known address", "The first priority rule sends property to the state of the owner's last-known address. Most items have no agent named. Without an address the property defaults to the holder entity's state of formation (second priority rule), not to the office's state.", "Agent master file / Paperless Pipeline; fill Detail cols S-T."),
    ("DATA GAP: legal entity and state of formation for each office", "Needed for the second-priority rule and to know which state's VDA to approach. 'Corporate HQ in Kansas City' is not the test - the formation state of each JV entity is.", "Legal / entity org chart."),
    ("CONFIRM: which tab is Florida and which is Alabama", "The analysis assumes Gallery = FL and Leading Edge = AL based on their banks. Both balances are small and entirely 2026, so the dollar effect of being wrong is nil today, but the rules differ (FL general period 5 yrs, AL 3 yrs).", "User."),
    ("CONFIRM: payable date vs. deposit date", "Dormancy runs from when the amount became payable to the owner, not from when the check was deposited. Deposit date is used as a proxy throughout. If agents are paid on closing, the closing date is the better trigger and could be earlier.", "Paperless Pipeline closing dates."),
    ("CONFIRM: independent-contractor status of agents (Texas especially)", "Texas' 1-year rule applies to 'wages' as defined in Labor Code 61.001, which excludes independent contractors. If agents are 1099 contractors, the 3-year general rule is the likely Texas test; if any are W-2, the 1-year rule applies to them.", "HR / agent agreements; confirm with counsel."),
    ("CONFIRM: prior-year write-offs", "Ask whether any older unapplied deposits were ever written off to revenue in prior years. Amounts owed to others that were taken into income are still reportable ('private escheat' is not permitted) and audits typically look back 10+ years.", "Accounting."),
    ("ACTION: work the oldest items first", "Houston has four items from 2022 and Dallas/TUR/Philly have 2023 items. Under the Texas 3-year test, every Texas item deposited before 3/1/2023 would already have been due on the July 1, 2026 report if it was owed to someone else; the Detail tab's 'Dormant date - LONG' column identifies them.", "Accounting."),
    ("ACTION: run due-diligence letters before any filing", "Every state in scope requires a written notice to the owner before remittance (TX: >$250, 60+ days before delivery; PA/IL/FL/AL/MO: $50 threshold per sources). A response from the owner stops the clock and lets you pay them directly instead of the state.", "Accounting; template letters."),
    ("ACTION: consider a Voluntary Disclosure Agreement for Texas", "Three Texas entities hold about 85% of the aged balance. Texas waives penalty and interest under a VDA but it must be requested before any Comptroller inquiry. Florida has a similar program. Confirm eligibility with the Comptroller or counsel.", "Counsel / Comptroller."),
    ("ACTION: items that are not the company's money", "Memo-flagged items (e.g. 'not for this address', 'Not in PP', 'duplicate?', 'Dallas transaction' deposited in Houston) are owed back to the payer, not to an agent. Return them to the payer now; a refund resolves the item and removes it from the escheat population.", "Accounting."),
    ("ACTION: legal hold item", "Houston row 27 ($1,500, 7/18/2023, memo: on hold per law enforcement). Keep segregated; document the hold; do not escheat or release without counsel.", "Counsel."),
    ("NOT VERIFIED in this session", "Statute and state-treasury web pages were blocked by this session's network policy, so every rule on the State Rules tab rests on search-engine summaries of those pages and on secondary sources (Sovos, UPPO, law-firm alerts). Specific items I could not confirm: Illinois business-holder report deadline; PA December 31 cutoff and VDA terms; FL and AL due-diligence thresholds; MO VDA. Treat the State Rules tab as a starting point for counsel, not a filing position.", "Counsel / unclaimed-property advisor."),
]
for i, (a, b, c) in enumerate(gaps, 4):
    put(gp, f"A{i}", i - 3); put(gp, f"B{i}", a, f_bold, wrap=True); put(gp, f"C{i}", b, wrap=True); put(gp, f"D{i}", c, wrap=True)
    gp.row_dimensions[i].height = 75
    for cc in range(1, 5): gp.cell(i, cc).border = box

# ================================================================ Method
mt = wb.create_sheet("Method", 4)
put(mt, "A1", "Method", f_title)
lines = [
    "1. Source file: United_Cash_Requirements.xlsx as provided on 10/6/2026 (file name f519c0da-United_Cash_Requirements.xlsx).",
    "2. For each entity tab, the 'Previous Month's Accrued Commissions' block was read (Dallas M20:R143, Chicago R20:V190, Houston M20:R286, Philly S21:W133, DC R22:V194, TUR M21:R299, Leading Edge AA22:AE195 and AG22:AK195, Gallery M21:R299). Rows without an 'x' in the 'accounted for in Sage' column are the unbooked items that make up the accrued balance.",
    "3. Each entity's extracted unbooked total was tied to the tab's 'outstanding amount' cell (Dallas M18 = 59,991.62; Houston M18 = 76,889.50; Philly S19 = 29,708.78; TUR M19 = 55,528.45; Leading Edge AA20 = 250.00 and AG20 = 395.00; Gallery M19 = 4,001.11; Chicago R18 = 0; DC R20 = 0). All tie exactly.",
    "4. The current-month 'Deposited checks' blocks were read the same way and tie to each tab's A18/A19 (and the second-bank cells) and to the source Summary row 5.",
    "5. URE Corp carries no accrued commissions (A4 = -O58 where O58 is empty).",
    "6. Ages are computed from the deposit date to the as-of date. Dormant dates = deposit date + the state period in months (EDATE).",
    "7. Classification hints are heuristics on amount and memo text only; they are not conclusions.",
    "8. State rules were researched by web search on 10/6/2026. Primary statute pages could not be opened from this session; see the State Rules tab 'Verification status' column for what is and is not confirmed.",
]
for i, t in enumerate(lines, 3):
    put(mt, f"A{i}", t, wrap=True); mt.merge_cells(f"A{i}:J{i}"); mt.row_dimensions[i].height = 34
mt.column_dimensions["A"].width = 120

wb.save(OUT)
print("saved", OUT, "detail rows", len(rows), "current rows", len(cur))
